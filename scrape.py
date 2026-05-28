import os
import json
import re
import requests
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("openpyxl not installed. Run: pip install openpyxl")

OUTPUT_FILE = "ecms_content.json"
DOCUMENTS_FOLDER = "./documents"
URLS_FILE = "./urls.txt"

TRADER_PATTERNS = [
    r'Exemption', 
    r'Duty free', 
    r'Annual Exemption',
    r'^1\.', r'^2\.1', r'^2\.2', r'^2\.3', r'^3\.1', r'^4\.', 
    r'^Declaration', r'^FE Export', r'^PCA', r'^TIER'
]

CUSTOMS_PATTERNS = [
    r'^C\.', r'^D\.', r'^E\.', r'^F\.', r'^G\.', r'^H\.', r'^i\.', r'^I\.',
    r'^ii\.', r'^J\.', r'^K\.', r'^L\.', r'^M\.'
]

def detect_side(filename):
    filename_lower = filename.lower()
    
    if any(x in filename_lower for x in ['exemption', 'duty free', 'refund', 'quota']):
        return 'TRADER'
        
    if any(x in filename_lower for x in ['registration', 'declaration', 'payment', 'pca', 'application']):
        return 'TRADER'
        
    if any(x in filename_lower for x in ['clearance', 'offence', 'risk', 'valuation', 'admin', 'system']):
        return 'CUSTOMS'

    for pattern in TRADER_PATTERNS:
        if re.search(pattern, filename, re.IGNORECASE):
            return 'TRADER'
    for pattern in CUSTOMS_PATTERNS:
        if re.search(pattern, filename, re.IGNORECASE):
            return 'CUSTOMS'
            
    return 'GENERAL'

def clean_filename(filename):
    name = Path(filename).stem
    name = re.sub(r'^\d+(\.\d+)?\s*[\.\-]?\s*', '', name)
    name = re.sub(r'^[A-Za-z]\.\s*', '', name)
    name = re.sub(r'^ii\.', '', name)
    name = re.sub(r'^FE\s+', '', name)
    return name.strip()

def extract_docx_enhanced(path):
    doc = Document(path)
    sections = []
    current_section = {"heading": "General", "content": []}
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        style_name = para.style.name.lower() if para.style else ""
        is_heading = (
            'heading' in style_name or 
            'title' in style_name or
            (len(text) < 100 and not text.endswith('.') and text[0].isupper() and not text[0].isdigit())
        )
        
        if is_heading and current_section["content"]:
            sections.append(current_section)
            current_section = {"heading": text, "content": []}
        else:
            current_section["content"].append(text)
    
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            cells = [cell.text.strip() if cell.text else "" for cell in row.cells]
            if any(cells):
                table_rows.append(" | ".join(c for c in cells if c))
        
        if table_rows:
            current_section["content"].append("[TABLE]\n" + "\n".join(table_rows))
    
    if current_section["content"]:
        sections.append(current_section)
    
    full_text = ""
    for sec in sections:
        content = "\n".join(sec["content"])
        full_text += f"\n\n## {sec['heading']}\n{content}"
    
    return full_text.strip()

def extract_excel_btc(path):
    if not HAS_OPENPYXL:
        return None
    
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        
        all_records = []
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            headers = []
            first_row = next(ws.iter_rows(values_only=True))
            for cell in first_row:
                headers.append(str(cell).strip().lower().replace(' ', '_') if cell else f"col_{len(headers)}")
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                
                record = {}
                for i, cell in enumerate(row):
                    if i < len(headers):
                        key = headers[i]
                        value = str(cell).strip() if cell is not None else ""
                        if value and value.lower() not in ['none', 'nan', 'null']:
                            record[key] = value
                
                if record:
                    all_records.append(record)
        
        if not all_records:
            return None
        
        content_parts = []
        content_parts.append(f"## BTC/HS Code Reference Database")
        content_parts.append(f"Total Records: {len(all_records)}")
        content_parts.append(f"Source File: {Path(path).name}")
        content_parts.append("")
        content_parts.append("## Search Guide")
        content_parts.append("Search by: code number, product description, or common name")
        content_parts.append("BTC codes are 8 digits (e.g., 01012100)")
        content_parts.append("HS codes are 6-10 digits (e.g., 0101.21)")
        content_parts.append("")
        
        for i, record in enumerate(all_records, 1):
            code = record.get('hs/btc_code', record.get('hs_code', record.get('btc_code', record.get('code', ''))))
            desc = record.get('description', '')
            common = record.get('common_name', '')
            
            if not code and not desc:
                continue
            
            entry_lines = [f"Entry {i}:"]
            entry_lines.append(f"  Code: {code}")
            entry_lines.append(f"  Description: {desc}")
            if common:
                entry_lines.append(f"  Common Names: {common}")
            
            keywords = []
            if desc:
                keywords.extend(re.findall(r'\b\w{3,}\b', desc.lower()))
            if common:
                keywords.extend(re.findall(r'\b\w{3,}\b', common.lower()))
            if code:
                keywords.append(code)
                if len(code) >= 6:
                    keywords.append(code[:6])
                    keywords.append(code[:4])
                    keywords.append(code[:2])
            
            if keywords:
                entry_lines.append(f"  Search Keywords: {', '.join(set(keywords))}")
            
            content_parts.append("\n".join(entry_lines))
        
        chapters = {}
        for record in all_records:
            code = record.get('hs/btc_code', record.get('hs_code', record.get('btc_code', '')))
            if len(str(code)) >= 2:
                chapter = str(code)[:2]
                if chapter not in chapters:
                    chapters[chapter] = []
                desc = record.get('description', '')[:50]
                chapters[chapter].append(f"{code}: {desc}")
        
        content_parts.append("")
        content_parts.append("## Chapter Summaries")
        for chapter, items in sorted(chapters.items())[:20]:
            content_parts.append(f"Chapter {chapter}: {len(items)} items")
            content_parts.append(f"  Examples: {'; '.join(items[:3])}")
        
        return "\n\n".join(content_parts)
        
    except Exception as e:
        print(f"Excel extraction failed: {e}")
        return None

def scrape_url(url):
    try:
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.5',
            'Connection': 'keep-alive',
        }
        
        print(f"Fetching {url}...")
        resp = requests.get(url, headers=headers, timeout=20, allow_redirects=True)
        
        if resp.url != url:
            print(f"Redirected to {resp.url}")
        
        resp.raise_for_status()
        
        content_type = resp.headers.get('content-type', '').lower()
        if 'text/html' not in content_type:
            print(f"Non-HTML content: {content_type}")
            return None
        
        soup = BeautifulSoup(resp.text, 'html.parser')
        
        for tag in soup(['script', 'style', 'nav', 'footer']):
            tag.decompose()
        
        title = soup.title.string.strip() if soup.title else url
        
        url_lower = url.lower()
        is_contact_page = any(x in url_lower for x in ['contact', 'focal', 'helpdesk', 'help-desk', 'support'])
        
        if not is_contact_page and title:
            title_lower = title.lower()
            is_contact_page = any(x in title_lower for x in ['contact', 'focal person', 'help desk', 'support'])
        
        if is_contact_page:
            print(f"Detected as CONTACT page")
            
            tables = soup.find_all('table')
            contact_text = f"SOURCE: {url}\nTITLE: {title}\nTYPE: CONTACT_PAGE\n\n"
            
            for table in tables:
                rows = table.find_all('tr')
                for row in rows:
                    cells = row.find_all(['td', 'th'])
                    if cells:
                        row_text = " | ".join(cell.get_text(strip=True) for cell in cells)
                        contact_text += row_text + "\n"
                contact_text += "\n"
            
            main = soup.find('main') or soup.find('article') or soup.find('div', class_=re.compile('content|main|body'))
            if main:
                text_elements = main.find_all(['p', 'div', 'li'])
                extra_lines = []
                for elem in text_elements:
                    text = elem.get_text(strip=True)
                    if text and len(text) > 10:
                        if any(indicator in text.lower() for indicator in ['phone', 'tel', 'email', '@', 'thimphu', 'gelephu', 'paro', 'samdrup', 'samtse', 'phuntsholing', 'kolkata', 'office', 'department']):
                            extra_lines.append(text)
                
                if extra_lines:
                    contact_text += "\nEXTRA CONTACT INFO:\n" + "\n".join(extra_lines[:50])
            
            if not tables:
                text = soup.get_text(separator='\n', strip=True)
                text = re.sub(r'\n+', '\n', text)
                contact_text += "\nPAGE CONTENT:\n" + text[:10000]
            
            return {
                "source": "contact_webpage",
                "title": "Contact Us",
                "content": contact_text[:30000],
                "type": "url",
                "side": "GENERAL"
            }
        
        main = soup.find('main') or soup.find('article') or soup.find('div', class_=re.compile('content|main'))
        text = main.get_text(separator='\n', strip=True) if main else soup.get_text(separator='\n', strip=True)
        
        text = re.sub(r'\n+', '\n', text)
        text = re.sub(r'\s+', ' ', text)
        
        return {
            "source": url,
            "title": title,
            "content": text[:15000],
            "type": "url",
            "side": "GENERAL"
        }
        
    except requests.exceptions.RequestException as e:
        print(f"Network error: {url} — {e}")
        return None
    except Exception as e:
        print(f"Failed: {url} — {e}")
        return None

def main():
    all_data = []
    
    doc_folder = Path(DOCUMENTS_FOLDER)
    if doc_folder.exists():
        docx_files = sorted(doc_folder.glob("*.docx"))
        print(f"Found {len(docx_files)} DOCX files")
        
        for file_path in docx_files:
            side = detect_side(file_path.name)
            clean_name = clean_filename(file_path.name)
            
            print(f"[{side}] {file_path.name} -> '{clean_name}'")
            
            text = extract_docx_enhanced(file_path)
            text = re.sub(r'\s+', ' ', text).strip()
            
            if text and len(text) > 50:
                contextual_text = f"[{side}] {clean_name}\n\n{text}"
                
                all_data.append({
                    "source": f"{side}_{clean_name}",
                    "title": clean_name,
                    "content": contextual_text,
                    "type": "document",
                    "side": side
                })
                print(f"{len(text)} chars")
            else:
                print("Empty or too short")
        
        if HAS_OPENPYXL:
            excel_files = sorted(doc_folder.glob("*.xlsx")) + sorted(doc_folder.glob("*.xls"))
            print(f"\nFound {len(excel_files)} Excel files")
            
            for file_path in excel_files:
                print(f"[BTC/HS] {file_path.name}")
                
                text = extract_excel_btc(file_path)
                
                if text and len(text) > 50:
                    source_name = f"BTC_{Path(file_path.name).stem}"
                    
                    all_data.append({
                        "source": source_name,
                        "title": f"BTC/HS Code Reference - {Path(file_path.name).stem}",
                        "content": f"[TARIFF] {source_name}\n\n{text}",
                        "type": "excel",
                        "side": "GENERAL"
                    })
                    print(f"{len(text)} chars")
                else:
                    print("Empty or failed to extract")
        else:
            print(f"\nopenpyxl not available, skipping Excel files")
    
    url_file = Path(URLS_FILE)
    if url_file.exists():
        urls = [u.strip() for u in url_file.read_text().splitlines() 
                if u.strip() and not u.strip().startswith('#')]
        
        print(f"\nFound {len(urls)} URLs")
        
        for url in urls:
            print(f"{url}")
            result = scrape_url(url)
            if result:
                all_data.append(result)
                print(f"Source: {result['source']} | {len(result['content'])} chars")
            else:
                print("Skipped (failed or empty)")
    else:
        print(f"\nURLs file not found: {URLS_FILE}")
    
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(all_data, f, ensure_ascii=False, indent=2)
    
    print(f"\nSaved {len(all_data)} sources to {OUTPUT_FILE}")
    print(f"Trader: {sum(1 for d in all_data if d.get('side') == 'TRADER')}")
    print(f"Customs: {sum(1 for d in all_data if d.get('side') == 'CUSTOMS')}")
    print(f"General: {sum(1 for d in all_data if d.get('side') == 'GENERAL')}")
    
    contact_sources = [d['source'] for d in all_data if 'contact' in d['source'].lower()]
    btc_sources = [d['source'] for d in all_data if d.get('type') == 'excel']
    
    if contact_sources:
        print(f"Contact sources: {contact_sources}")
    else:
        print("No contact sources found!")
    
    if btc_sources:
        print(f"BTC/HS sources: {btc_sources}")

if __name__ == "__main__":
    main()